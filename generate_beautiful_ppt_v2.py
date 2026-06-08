#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
智能对话伙伴 - PPT生成脚本
科技感+亲和力，蓝白配色，橙色点缀
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

# 配色定义
BLUE = RGBColor(30, 80, 150)      # 深蓝
LIGHT_BLUE = RGBColor(100, 150, 220)  # 浅蓝
ORANGE = RGBColor(255, 130, 50)   # 橙色
WHITE = RGBColor(255, 255, 255)
LIGHT_GRAY = RGBColor(245, 248, 252)
GRAY = RGBColor(120, 120, 120)

# 创建演示文稿
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def add_logo(slide):
    """添加统一logo在右下角"""
    # 两个重叠的对话气泡
    bubble1 = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(11.5), Inches(6.8),
        Inches(1.2), Inches(0.6)
    )
    bubble1.fill.solid()
    bubble1.fill.fore_color.rgb = BLUE
    bubble1.line.fill.background()
    
    bubble2 = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(11.7), Inches(6.9),
        Inches(1.2), Inches(0.6)
    )
    bubble2.fill.solid()
    bubble2.fill.fore_color.rgb = ORANGE
    bubble2.line.fill.background()

def add_cover_slide():
    """第1页：封面"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # 背景
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_GRAY
    bg.line.fill.background()
    
    # 装饰元素
    circle1 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1), Inches(1), Inches(4), Inches(4))
    circle1.fill.solid()
    circle1.fill.fore_color.rgb = BLUE
    circle1.line.fill.background()
    circle1.fill.transparency = 0.85
    
    circle2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8), Inches(2), Inches(3.5), Inches(3.5))
    circle2.fill.solid()
    circle2.fill.fore_color.rgb = ORANGE
    circle2.line.fill.background()
    circle2.fill.transparency = 0.85
    
    # 主标题
    title_box = slide.shapes.add_textbox(Inches(2), Inches(2.5), Inches(9), Inches(1.5))
    tf = title_box.text_frame
    tf.paragraphs[0].text = "智能对话伙伴"
    tf.paragraphs[0].font.size = Pt(54)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BLUE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # 副标题
    subtitle_box = slide.shapes.add_textbox(Inches(2), Inches(4), Inches(9), Inches(0.8))
    tf = subtitle_box.text_frame
    tf.paragraphs[0].text = "少儿口语练习智能体 × 二三级口译智能体"
    tf.paragraphs[0].font.size = Pt(26)
    tf.paragraphs[0].font.color.rgb = GRAY
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # 底部文字
    bottom_box = slide.shapes.add_textbox(Inches(2), Inches(5.8), Inches(9), Inches(0.6))
    tf = bottom_box.text_frame
    tf.paragraphs[0].text = "让语言学习更专业、更轻松"
    tf.paragraphs[0].font.size = Pt(22)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = ORANGE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    add_logo(slide)

def add_overview_slide():
    """第2页：总览页"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    bg.line.fill.background()
    
    # 标题
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    tf.paragraphs[0].text = "两大智能体，覆盖语言成长全阶段"
    tf.paragraphs[0].font.size = Pt(36)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BLUE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # 左侧卡片
    left_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.5), Inches(5), Inches(4.5))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = RGBColor(240, 250, 255)
    left_card.line.color.rgb = BLUE
    left_card.line.width = Pt(2)
    
    left_title = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(4.4), Inches(0.6))
    tf = left_title.text_frame
    tf.paragraphs[0].text = "少儿口语练习智能体"
    tf.paragraphs[0].font.size = Pt(22)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = ORANGE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    left_desc = slide.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(4.4), Inches(1))
    tf = left_desc.text_frame
    tf.paragraphs[0].text = "兴趣启蒙·自信开口"
    tf.paragraphs[0].font.size = Pt(18)
    tf.paragraphs[0].font.color.rgb = GRAY
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # 右侧卡片
    right_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), Inches(1.5), Inches(5), Inches(4.5))
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = RGBColor(245, 248, 255)
    right_card.line.color.rgb = BLUE
    right_card.line.width = Pt(2)
    
    right_title = slide.shapes.add_textbox(Inches(8.1), Inches(1.8), Inches(4.4), Inches(0.6))
    tf = right_title.text_frame
    tf.paragraphs[0].text = "二三级口译智能体"
    tf.paragraphs[0].font.size = Pt(22)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BLUE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    right_desc = slide.shapes.add_textbox(Inches(8.1), Inches(2.5), Inches(4.4), Inches(1))
    tf = right_desc.text_frame
    tf.paragraphs[0].text = "精准高效·专业进阶"
    tf.paragraphs[0].font.size = Pt(18)
    tf.paragraphs[0].font.color.rgb = GRAY
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # 中间箭头
    arrow_shape = slide.shapes.add_shape(MSO_SHAPE.UP_ARROW, Inches(6), Inches(3), Inches(1.3), Inches(2))
    arrow_shape.fill.solid()
    arrow_shape.fill.fore_color.rgb = ORANGE
    arrow_shape.line.fill.background()
    
    # 底部文字
    bottom_text = slide.shapes.add_textbox(Inches(1), Inches(6.3), Inches(11), Inches(0.5))
    tf = bottom_text.text_frame
    tf.paragraphs[0].text = "从敢说到会译，每一步都有AI陪伴。"
    tf.paragraphs[0].font.size = Pt(20)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = GRAY
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    add_logo(slide)

def add_kids_intro_slide():
    """第3页：智能体一介绍"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_GRAY
    bg.line.fill.background()
    
    # 标题
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    tf.paragraphs[0].text = "少儿口语练习智能体 — 让孩子爱上说英语"
    tf.paragraphs[0].font.size = Pt(34)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = ORANGE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # 三个功能卡片
    features = [
        ("场景对话", "模拟餐厅、动物园等生活场景"),
        ("语音打分", "实时纠正发音，像游戏一样闯关"),
        ("趣味激励", "动画勋章+成长树")
    ]
    
    for i, (title, desc) in enumerate(features):
        x = 0.8 + i * 4.2
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.8), Inches(3.8), Inches(3.5))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = ORANGE
        card.line.width = Pt(2)
        
        title_box = slide.shapes.add_textbox(Inches(x + 0.3), Inches(2), Inches(3), Inches(0.6))
        tf = title_box.text_frame
        tf.paragraphs[0].text = title
        tf.paragraphs[0].font.size = Pt(22)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = ORANGE
        
        desc_box = slide.shapes.add_textbox(Inches(x + 0.3), Inches(2.8), Inches(3), Inches(1.5))
        tf = desc_box.text_frame
        tf.paragraphs[0].text = desc
        tf.paragraphs[0].font.size = Pt(16)
        tf.paragraphs[0].font.color.rgb = GRAY
    
    # 底部装饰文字
    bottom_box = slide.shapes.add_textbox(Inches(1), Inches(5.7), Inches(11), Inches(0.6))
    tf = bottom_box.text_frame
    tf.paragraphs[0].text = "🎤  😊  🌱"
    tf.paragraphs[0].font.size = Pt(30)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    add_logo(slide)

def add_kids_scenarios_slide():
    """第4页：智能体一的应用场景"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    bg.line.fill.background()
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    tf.paragraphs[0].text = "轻松融入学习日常"
    tf.paragraphs[0].font.size = Pt(36)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BLUE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    scenarios = [
        ("家庭", "睡前10分钟互动练习"),
        ("课堂", "老师的小助手，分组PK"),
        ("自学", "孩子自己点开就能玩")
    ]
    
    for i, (title, desc) in enumerate(scenarios):
        x = 0.8 + i * 4.2
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.8), Inches(3.8), Inches(4))
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_GRAY
        card.line.color.rgb = BLUE
        card.line.width = Pt(2)
        
        t_box = slide.shapes.add_textbox(Inches(x + 0.3), Inches(2), Inches(3), Inches(0.6))
        tf = t_box.text_frame
        tf.paragraphs[0].text = title
        tf.paragraphs[0].font.size = Pt(24)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = ORANGE
        
        d_box = slide.shapes.add_textbox(Inches(x + 0.3), Inches(2.8), Inches(3), Inches(1.5))
        tf = d_box.text_frame
        tf.paragraphs[0].text = desc
        tf.paragraphs[0].font.size = Pt(16)
        tf.paragraphs[0].font.color.rgb = GRAY
    
    add_logo(slide)

def add_kids_workflow_slide():
    """第5页：智能体一的工作流"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_GRAY
    bg.line.fill.background()
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    tf.paragraphs[0].text = "一次开口，三步成长"
    tf.paragraphs[0].font.size = Pt(36)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = ORANGE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    steps = [
        ("选择场景", "孩子挑选主题\n（如“在动物园”）"),
        ("互动对话", "AI用卡通角色提问\n孩子语音回答"),
        ("反馈激励", "实时打分+收集勋章\n小树长高")
    ]
    
    step_colors = [ORANGE, RGBColor(240, 120, 50), RGBColor(255, 100, 30)]
    
    for i, (title, desc) in enumerate(steps):
        x = 1 + i * 4
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.5), Inches(3.6), Inches(3.5))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = step_colors[i]
        card.line.width = Pt(3)
        
        t_box = slide.shapes.add_textbox(Inches(x + 0.2), Inches(1.7), Inches(3.2), Inches(0.6))
        tf = t_box.text_frame
        tf.paragraphs[0].text = f"{i+1}. {title}"
        tf.paragraphs[0].font.size = Pt(20)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = step_colors[i]
        
        d_box = slide.shapes.add_textbox(Inches(x + 0.2), Inches(2.5), Inches(3.2), Inches(2))
        tf = d_box.text_frame
        tf.paragraphs[0].text = desc
        tf.paragraphs[0].font.size = Pt(14)
        tf.paragraphs[0].font.color.rgb = GRAY
    
    # 箭头连接
    arrow1 = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.2), Inches(3), Inches(0.8), Inches(0.5))
    arrow1.fill.solid()
    arrow1.fill.fore_color.rgb = GRAY
    arrow1.line.fill.background()
    
    arrow2 = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(8.2), Inches(3), Inches(0.8), Inches(0.5))
    arrow2.fill.solid()
    arrow2.fill.fore_color.rgb = GRAY
    arrow2.line.fill.background()
    
    bottom_text = slide.shapes.add_textbox(Inches(1), Inches(5.5), Inches(11), Inches(0.6))
    tf = bottom_text.text_frame
    tf.paragraphs[0].text = "每次3-5分钟，像玩游戏一样练口语"
    tf.paragraphs[0].font.size = Pt(20)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = ORANGE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    add_logo(slide)

def add_interpreter_intro_slide():
    """第6页：智能体二介绍"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    bg.line.fill.background()
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    tf.paragraphs[0].text = "二三级口译智能体 — 向专业口译迈进"
    tf.paragraphs[0].font.size = Pt(34)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BLUE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    features = [
        ("双向交传模拟", "中英互译，自动分段录音"),
        ("笔记法训练", "关键信息捕捉提示"),
        ("考试真题演练", "覆盖CATTI二三级题型")
    ]
    
    for i, (title, desc) in enumerate(features):
        x = 0.8 + i * 4.2
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.8), Inches(3.8), Inches(3.5))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(248, 250, 255)
        card.line.color.rgb = BLUE
        card.line.width = Pt(2)
        
        t_box = slide.shapes.add_textbox(Inches(x + 0.3), Inches(2), Inches(3), Inches(0.6))
        tf = t_box.text_frame
        tf.paragraphs[0].text = title
        tf.paragraphs[0].font.size = Pt(20)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = BLUE
        
        d_box = slide.shapes.add_textbox(Inches(x + 0.3), Inches(2.8), Inches(3), Inches(1.5))
        tf = d_box.text_frame
        tf.paragraphs[0].text = desc
        tf.paragraphs[0].font.size = Pt(16)
        tf.paragraphs[0].font.color.rgb = GRAY
    
    bottom_box = slide.shapes.add_textbox(Inches(1), Inches(5.7), Inches(11), Inches(0.6))
    tf = bottom_box.text_frame
    tf.paragraphs[0].text = "🎧  📝  🏆"
    tf.paragraphs[0].font.size = Pt(30)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    add_logo(slide)

def add_interpreter_scenarios_slide():
    """第7页：智能体二的应用场景"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_GRAY
    bg.line.fill.background()
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    tf.paragraphs[0].text = "备考与实战利器"
    tf.paragraphs[0].font.size = Pt(36)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BLUE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    scenarios = [
        ("备考自习", "自由设定语速与停顿时间"),
        ("课堂演练", "模拟会议口译，可复盘录音"),
        ("自我评估", "AI给出准确率与流畅度分析")
    ]
    
    for i, (title, desc) in enumerate(scenarios):
        x = 0.8 + i * 4.2
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.8), Inches(3.8), Inches(4))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = BLUE
        card.line.width = Pt(2)
        
        t_box = slide.shapes.add_textbox(Inches(x + 0.3), Inches(2), Inches(3), Inches(0.6))
        tf = t_box.text_frame
        tf.paragraphs[0].text = title
        tf.paragraphs[0].font.size = Pt(22)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = BLUE
        
        d_box = slide.shapes.add_textbox(Inches(x + 0.3), Inches(2.8), Inches(3), Inches(1.5))
        tf = d_box.text_frame
        tf.paragraphs[0].text = desc
        tf.paragraphs[0].font.size = Pt(16)
        tf.paragraphs[0].font.color.rgb = GRAY
    
    add_logo(slide)

def add_interpreter_workflow_slide():
    """第8页：智能体二的工作流"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    bg.line.fill.background()
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    tf.paragraphs[0].text = "一段录音，四步进阶"
    tf.paragraphs[0].font.size = Pt(36)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BLUE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    steps = [
        ("获取源语", "播放真题/用户上传音频\n（中/英）"),
        ("口译录音", "听到提示音后开始口译\n自动分段"),
        ("AI评估", "从准确率、流畅度、术语\n三方面打分"),
        ("复盘学习", "对比参考译文，标记薄弱句\n支持重练")
    ]
    
    for i, (title, desc) in enumerate(steps):
        x = 0.6 + i * 3.2
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.5), Inches(2.8), Inches(4))
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_GRAY
        card.line.color.rgb = BLUE
        card.line.width = Pt(2)
        
        t_box = slide.shapes.add_textbox(Inches(x + 0.15), Inches(1.7), Inches(2.5), Inches(0.6))
        tf = t_box.text_frame
        tf.paragraphs[0].text = f"{i+1}. {title}"
        tf.paragraphs[0].font.size = Pt(16)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = BLUE
        
        d_box = slide.shapes.add_textbox(Inches(x + 0.15), Inches(2.4), Inches(2.5), Inches(2.5))
        tf = d_box.text_frame
        tf.paragraphs[0].text = desc
        tf.paragraphs[0].font.size = Pt(12)
        tf.paragraphs[0].font.color.rgb = GRAY
    
    bottom_text = slide.shapes.add_textbox(Inches(1), Inches(5.8), Inches(11), Inches(0.6))
    tf = bottom_text.text_frame
    tf.paragraphs[0].text = "练一次，进步一次，随时回看成长曲线"
    tf.paragraphs[0].font.size = Pt(18)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BLUE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    add_logo(slide)

def add_comparison_slide():
    """第9页：对比与协作"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_GRAY
    bg.line.fill.background()
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    tf.paragraphs[0].text = "一个家庭，两种需要 —— 智能体如何一起使用？"
    tf.paragraphs[0].font.size = Pt(32)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BLUE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # 左侧
    left_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.5), Inches(5.5), Inches(4))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = WHITE
    left_card.line.color.rgb = ORANGE
    left_card.line.width = Pt(2)
    
    left_title = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(4.9), Inches(0.6))
    tf = left_title.text_frame
    tf.paragraphs[0].text = "少儿口语"
    tf.paragraphs[0].font.size = Pt(24)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = ORANGE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    left_desc = slide.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(4.9), Inches(2))
    tf = left_desc.text_frame
    tf.paragraphs[0].text = "亲子共学，培养语感\n激发兴趣，自信开口"
    tf.paragraphs[0].font.size = Pt(18)
    tf.paragraphs[0].font.color.rgb = GRAY
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # 中间连接
    center_box = slide.shapes.add_textbox(Inches(6.2), Inches(3), Inches(0.8), Inches(1))
    tf = center_box.text_frame
    tf.paragraphs[0].text = "➡️"
    tf.paragraphs[0].font.size = Pt(40)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # 右侧
    right_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.3), Inches(1.5), Inches(5.5), Inches(4))
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = WHITE
    right_card.line.color.rgb = BLUE
    right_card.line.width = Pt(2)
    
    right_title = slide.shapes.add_textbox(Inches(7.6), Inches(1.8), Inches(4.9), Inches(0.6))
    tf = right_title.text_frame
    tf.paragraphs[0].text = "二三级口译"
    tf.paragraphs[0].font.size = Pt(24)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BLUE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    right_desc = slide.shapes.add_textbox(Inches(7.6), Inches(2.6), Inches(4.9), Inches(2))
    tf = right_desc.text_frame
    tf.paragraphs[0].text = "个人精进，职业跳板\n专业训练，高效备考"
    tf.paragraphs[0].font.size = Pt(18)
    tf.paragraphs[0].font.color.rgb = GRAY
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # 中间连接语
    conn_text = slide.shapes.add_textbox(Inches(3), Inches(3.8), Inches(7), Inches(0.6))
    tf = conn_text.text_frame
    tf.paragraphs[0].text = "从少儿到成人的语言能力进阶之路"
    tf.paragraphs[0].font.size = Pt(18)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = GRAY
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    bottom_text = slide.shapes.add_textbox(Inches(1), Inches(5.8), Inches(11), Inches(0.6))
    tf = bottom_text.text_frame
    tf.paragraphs[0].text = "同一个账号，记录全周期成长轨迹。"
    tf.paragraphs[0].font.size = Pt(20)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BLUE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    add_logo(slide)

def add_summary_slide():
    """第10页：总结与展望"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    bg.line.fill.background()
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    tf.paragraphs[0].text = "让AI成为语言成长的终身伙伴"
    tf.paragraphs[0].font.size = Pt(38)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BLUE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    points = [
        ("个性化", "匹配不同年龄与水平"),
        ("便捷性", "手机/平板/电脑随时用"),
        ("未来计划", "增加手语翻译、方言口音适应")
    ]
    
    point_colors = [ORANGE, BLUE, RGBColor(80, 160, 120)]
    
    for i, (title, desc) in enumerate(points):
        y = 1.8 + i * 1.5
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.5), Inches(y), Inches(8), Inches(1.2))
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT_GRAY
        card.line.color.rgb = point_colors[i]
        card.line.width = Pt(2)
        
        t_box = slide.shapes.add_textbox(Inches(2.8), Inches(y + 0.15), Inches(7.4), Inches(0.5))
        tf = t_box.text_frame
        tf.paragraphs[0].text = title
        tf.paragraphs[0].font.size = Pt(22)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = point_colors[i]
        
        d_box = slide.shapes.add_textbox(Inches(2.8), Inches(y + 0.65), Inches(7.4), Inches(0.4))
        tf = d_box.text_frame
        tf.paragraphs[0].text = desc
        tf.paragraphs[0].font.size = Pt(16)
        tf.paragraphs[0].font.color.rgb = GRAY
    
    bottom_slogan = slide.shapes.add_textbox(Inches(1), Inches(6.2), Inches(11), Inches(0.6))
    tf = bottom_slogan.text_frame
    tf.paragraphs[0].text = "语言无界，成长可见"
    tf.paragraphs[0].font.size = Pt(28)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = ORANGE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    add_logo(slide)

def add_thanks_slide():
    """第11页：致谢/结束页"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_GRAY
    bg.line.fill.background()
    
    # 装饰背景
    circle1 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.5), Inches(1.5), Inches(4), Inches(4))
    circle1.fill.solid()
    circle1.fill.fore_color.rgb = ORANGE
    circle1.line.fill.background()
    circle1.fill.transparency = 0.9
    
    circle2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.8), Inches(1.5), Inches(4), Inches(4))
    circle2.fill.solid()
    circle2.fill.fore_color.rgb = BLUE
    circle2.line.fill.background()
    circle2.fill.transparency = 0.9
    
    # 主标题
    title_box = slide.shapes.add_textbox(Inches(2), Inches(2), Inches(9), Inches(1.5))
    tf = title_box.text_frame
    tf.paragraphs[0].text = "感谢聆听"
    tf.paragraphs[0].font.size = Pt(56)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BLUE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # 副标题
    subtitle_box = slide.shapes.add_textbox(Inches(2), Inches(3.8), Inches(9), Inches(0.8))
    tf = subtitle_box.text_frame
    tf.paragraphs[0].text = "智能对话伙伴，陪伴每一次成长"
    tf.paragraphs[0].font.size = Pt(26)
    tf.paragraphs[0].font.color.rgb = GRAY
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # Q&A
    qa_box = slide.shapes.add_textbox(Inches(2), Inches(5.5), Inches(9), Inches(0.8))
    tf = qa_box.text_frame
    tf.paragraphs[0].text = "Q&A"
    tf.paragraphs[0].font.size = Pt(32)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = ORANGE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    add_logo(slide)

# 生成所有页面
add_cover_slide()
add_overview_slide()
add_kids_intro_slide()
add_kids_scenarios_slide()
add_kids_workflow_slide()
add_interpreter_intro_slide()
add_interpreter_scenarios_slide()
add_interpreter_workflow_slide()
add_comparison_slide()
add_summary_slide()
add_thanks_slide()

output_path = "智能对话伙伴_最终版.pptx"
prs.save(output_path)
print(f"✅ 精美PPT已生成: {output_path}")
