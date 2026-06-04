#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
AI赋能语言进阶 - 最终版PPT生成脚本
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# 配色定义
ORANGE = RGBColor(255, 107, 53)    # FF6B35
BLUE = RGBColor(30, 58, 95)        # 1E3A5F
LIGHT_GRAY = RGBColor(245, 247, 250)
DARK_GRAY = RGBColor(71, 85, 105)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(0, 0, 0)

# 创建演示文稿
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def add_cover_slide():
    """第1页：封面"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_GRAY
    bg.line.fill.background()
    
    left_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(6.666), prs.slide_height)
    left_bg.fill.solid()
    left_bg.fill.fore_color.rgb = ORANGE
    left_bg.line.fill.background()
    
    right_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.666), 0, Inches(6.666), prs.slide_height)
    right_bg.fill.solid()
    right_bg.fill.fore_color.rgb = BLUE
    right_bg.line.fill.background()
    
    title_box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(11), Inches(1.5))
    tf = title_box.text_frame
    tf.paragraphs[0].text = "AI 赋能语言进阶"
    tf.paragraphs[0].font.size = Pt(56)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    subtitle_box = slide.shapes.add_textbox(Inches(1), Inches(3.8), Inches(11), Inches(1))
    tf = subtitle_box.text_frame
    tf.paragraphs[0].text = "少儿口语练习智能体 × 三级口译智能体"
    tf.paragraphs[0].font.size = Pt(28)
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    footer_box = slide.shapes.add_textbox(Inches(1), Inches(6.2), Inches(11), Inches(0.8))
    tf = footer_box.text_frame
    tf.paragraphs[0].text = "语言学习全链路解决方案 | AI团队 | 2026.06"
    tf.paragraphs[0].font.size = Pt(16)
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

def add_overview_slide():
    """第2页：全景概览"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_GRAY
    bg.line.fill.background()
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    tf.paragraphs[0].text = "两个智能体，覆盖语言学习的两大阶段"
    tf.paragraphs[0].font.size = Pt(36)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BLUE
    
    left_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.5), Inches(5.8), Inches(5.5))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = ORANGE
    left_card.line.fill.background()
    
    right_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7), Inches(1.5), Inches(5.8), Inches(5.5))
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = BLUE
    right_card.line.fill.background()
    
    left_title = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.2), Inches(0.6))
    tf = left_title.text_frame
    tf.paragraphs[0].text = "少儿口语练习智能体"
    tf.paragraphs[0].font.size = Pt(22)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    right_title = slide.shapes.add_textbox(Inches(7.3), Inches(1.8), Inches(5.2), Inches(0.6))
    tf = right_title.text_frame
    tf.paragraphs[0].text = "三级口译智能体"
    tf.paragraphs[0].font.size = Pt(22)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    left_content = [
        "核心定位: 兴趣启蒙 · 自信开口",
        "目标用户: 4–12 岁少儿",
        "核心能力: 情景对话 + 语音评分 + 游戏化",
        "输出形式: 鼓励性反馈 + 动画奖励"
    ]
    y = 2.7
    for line in left_content:
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(y), Inches(5.2), Inches(0.6))
        tf = tb.text_frame
        tf.paragraphs[0].text = line
        tf.paragraphs[0].font.size = Pt(14)
        tf.paragraphs[0].font.color.rgb = WHITE
        y += 0.7
    
    right_content = [
        "核心定位: 技能进阶 · 考场实战",
        "目标用户: 备考CATTI三级口译的学习者",
        "核心能力: 交传训练 + 数字复述 + 高频真题",
        "输出形式: 精准打分 + 慢速/常速示范"
    ]
    y = 2.7
    for line in right_content:
        tb = slide.shapes.add_textbox(Inches(7.3), Inches(y), Inches(5.2), Inches(0.6))
        tf = tb.text_frame
        tf.paragraphs[0].text = line
        tf.paragraphs[0].font.size = Pt(14)
        tf.paragraphs[0].font.color.rgb = WHITE
        y += 0.7

def add_kids_features_slide():
    """第3页：少儿口语智能体 – 功能"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_GRAY
    bg.line.fill.background()
    
    title_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1))
    title_bg.fill.solid()
    title_bg.fill.fore_color.rgb = ORANGE
    title_bg.line.fill.background()
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12), Inches(0.7))
    tf = title_box.text_frame
    tf.paragraphs[0].text = "少儿口语练习智能体"
    tf.paragraphs[0].font.size = Pt(32)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(1), Inches(12), Inches(0.5))
    tf = subtitle_box.text_frame
    tf.paragraphs[0].text = "让孩子像玩游戏一样爱上开口"
    tf.paragraphs[0].font.size = Pt(20)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = DARK_GRAY
    
    features = [
        ("🎙️", "AI 情景对话", "学校/动物园/生日派对等12个主题场景"),
        ("⭐", "儿童语音识别", "专为不标准发音优化，宽容度更高"),
        ("🎮", "跟读打卡", "每完成一关获得星星勋章"),
        ("📊", "成长报告", "记录每周开口时长 & 进步词汇量"),
    ]
    
    for i, (icon, title, desc) in enumerate(features):
        x = 0.7 + (i % 2) * 6
        y = 1.8 + (i // 2) * 2.8
        
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.5), Inches(2.4))
        card.fill.solid()
        card.fill.fore_color.rgb = ORANGE
        card.line.fill.background()
        
        icon_box = slide.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.2), Inches(1), Inches(0.6))
        tf = icon_box.text_frame
        tf.paragraphs[0].text = icon
        tf.paragraphs[0].font.size = Pt(28)
        
        title_box = slide.shapes.add_textbox(Inches(x + 1.2), Inches(y + 0.2), Inches(4), Inches(0.5))
        tf = title_box.text_frame
        tf.paragraphs[0].text = title
        tf.paragraphs[0].font.size = Pt(18)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = WHITE
        
        desc_box = slide.shapes.add_textbox(Inches(x + 0.3), Inches(y + 0.9), Inches(5), Inches(1.3))
        tf = desc_box.text_frame
        tf.paragraphs[0].text = desc
        tf.paragraphs[0].font.size = Pt(14)
        tf.paragraphs[0].font.color.rgb = WHITE
    
    example_box = slide.shapes.add_textbox(Inches(0.5), Inches(6), Inches(12), Inches(0.8))
    tf = example_box.text_frame
    tf.paragraphs[0].text = '"老师好，我昨天去动物园了" → 智能体模拟动物园管理员回应并追问："你最喜欢什么动物呀？"'
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.color.rgb = DARK_GRAY
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

def add_kids_scenarios_slide():
    """第4页：少儿口语智能体 – 特色与场景"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_GRAY
    bg.line.fill.background()
    
    title_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1))
    title_bg.fill.solid()
    title_bg.fill.fore_color.rgb = ORANGE
    title_bg.line.fill.background()
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12), Inches(0.7))
    tf = title_box.text_frame
    tf.paragraphs[0].text = "不只是“练口语”，更是“敢表达”"
    tf.paragraphs[0].font.size = Pt(32)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    
    left_section = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.3), Inches(5.8), Inches(5.5))
    left_section.fill.solid()
    left_section.fill.fore_color.rgb = WHITE
    left_section.line.fill.background()
    
    left_title = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5), Inches(0.5))
    tf = left_title.text_frame
    tf.paragraphs[0].text = "两大特色"
    tf.paragraphs[0].font.size = Pt(22)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = ORANGE
    
    features = [
        "1. 无压力纠错：不打断孩子，对话结束后用示范句自然重复正确表达",
        "2. 亲子模式：家长可查看对话原文 + 建议互动话题"
    ]
    y = 2.2
    for f in features:
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(y), Inches(5), Inches(1))
        tf = tb.text_frame
        tf.paragraphs[0].text = f
        tf.paragraphs[0].font.size = Pt(14)
        tf.paragraphs[0].font.color.rgb = DARK_GRAY
        y += 1.3
    
    right_section = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7), Inches(1.3), Inches(5.8), Inches(5.5))
    right_section.fill.solid()
    right_section.fill.fore_color.rgb = WHITE
    right_section.line.fill.background()
    
    right_title = slide.shapes.add_textbox(Inches(7.3), Inches(1.5), Inches(5), Inches(0.5))
    tf = right_title.text_frame
    tf.paragraphs[0].text = "典型使用场景"
    tf.paragraphs[0].font.size = Pt(22)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = ORANGE
    
    scenarios = [
        "🕐 课前5分钟：与智能体预热，降低课堂紧张感",
        "👨 课后10分钟：角色扮演作业伙伴，替代死记硬背",
        "🚗 碎片时间：车程中完成一次“动物园导览”口语挑战"
    ]
    y = 2.2
    for s in scenarios:
        tb = slide.shapes.add_textbox(Inches(7.3), Inches(y), Inches(5), Inches(0.8))
        tf = tb.text_frame
        tf.paragraphs[0].text = s
        tf.paragraphs[0].font.size = Pt(14)
        tf.paragraphs[0].font.color.rgb = DARK_GRAY
        y += 1.1

def add_interpreter_features_slide():
    """第5页：三级口译智能体 – 功能"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_GRAY
    bg.line.fill.background()
    
    title_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1))
    title_bg.fill.solid()
    title_bg.fill.fore_color.rgb = BLUE
    title_bg.line.fill.background()
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12), Inches(0.7))
    tf = title_box.text_frame
    tf.paragraphs[0].text = "三级口译智能体"
    tf.paragraphs[0].font.size = Pt(32)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(1), Inches(12), Inches(0.5))
    tf = subtitle_box.text_frame
    tf.paragraphs[0].text = "交传训练的专业级AI陪练"
    tf.paragraphs[0].font.size = Pt(20)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = DARK_GRAY
    
    features = [
        ("🧠", "英汉 / 汉英双向交传", "30秒–1分钟分段练习"),
        ("📝", "数字复述强化", "专门训练数字、日期、百分比的准确转译"),
        ("🎧", "语速自适应", "可从80词/分钟调至140词/分钟"),
        ("✅", "评分维度", "信息完整度40% + 表达流畅度30% + 术语准确度30%"),
    ]
    
    for i, (icon, title, desc) in enumerate(features):
        x = 0.7 + (i % 2) * 6
        y = 1.8 + (i // 2) * 2.8
        
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.5), Inches(2.4))
        card.fill.solid()
        card.fill.fore_color.rgb = BLUE
        card.line.fill.background()
        
        icon_box = slide.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.2), Inches(1), Inches(0.6))
        tf = icon_box.text_frame
        tf.paragraphs[0].text = icon
        tf.paragraphs[0].font.size = Pt(28)
        
        title_box = slide.shapes.add_textbox(Inches(x + 1.2), Inches(y + 0.2), Inches(4), Inches(0.5))
        tf = title_box.text_frame
        tf.paragraphs[0].text = title
        tf.paragraphs[0].font.size = Pt(18)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = WHITE
        
        desc_box = slide.shapes.add_textbox(Inches(x + 0.3), Inches(y + 0.9), Inches(5), Inches(1.3))
        tf = desc_box.text_frame
        tf.paragraphs[0].text = desc
        tf.paragraphs[0].font.size = Pt(14)
        tf.paragraphs[0].font.color.rgb = WHITE
    
    data_box = slide.shapes.add_textbox(Inches(0.5), Inches(6), Inches(12), Inches(0.8))
    tf = data_box.text_frame
    tf.paragraphs[0].text = "📈 某用户练习2周后：数字准确率从62%提升至89%"
    tf.paragraphs[0].font.size = Pt(18)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BLUE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

def add_interpreter_scenarios_slide():
    """第6页：三级口译智能体 – 特色与场景"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_GRAY
    bg.line.fill.background()
    
    title_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1))
    title_bg.fill.solid()
    title_bg.fill.fore_color.rgb = BLUE
    title_bg.line.fill.background()
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12), Inches(0.7))
    tf = title_box.text_frame
    tf.paragraphs[0].text = "从“听懂了”到“译得出”"
    tf.paragraphs[0].font.size = Pt(32)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    
    left_section = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.3), Inches(5.8), Inches(5.5))
    left_section.fill.solid()
    left_section.fill.fore_color.rgb = WHITE
    left_section.line.fill.background()
    
    left_title = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5), Inches(0.5))
    tf = left_title.text_frame
    tf.paragraphs[0].text = "两大特色"
    tf.paragraphs[0].font.size = Pt(22)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BLUE
    
    features = [
        "1. 影子跟读 + 带稿同传双模式：过渡训练，降低认知负荷",
        "2. 错误聚类分析：智能标注常错术语（如multilateral连续误听为multiple）"
    ]
    y = 2.2
    for f in features:
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(y), Inches(5), Inches(1))
        tf = tb.text_frame
        tf.paragraphs[0].text = f
        tf.paragraphs[0].font.size = Pt(14)
        tf.paragraphs[0].font.color.rgb = DARK_GRAY
        y += 1.3
    
    right_section = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7), Inches(1.3), Inches(5.8), Inches(5.5))
    right_section.fill.solid()
    right_section.fill.fore_color.rgb = WHITE
    right_section.line.fill.background()
    
    right_title = slide.shapes.add_textbox(Inches(7.3), Inches(1.5), Inches(5), Inches(0.5))
    tf = right_title.text_frame
    tf.paragraphs[0].text = "典型使用场景"
    tf.paragraphs[0].font.size = Pt(22)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BLUE
    
    scenarios = [
        "🌅 每日30分钟：晨间数字复述（专攻CATTI最易失分项）",
        "📚 真题模拟：200+道历年三级口译真题情景复现",
        "⏰ 倒计时模式：训练考场心理素质与笔记法"
    ]
    y = 2.2
    for s in scenarios:
        tb = slide.shapes.add_textbox(Inches(7.3), Inches(y), Inches(5), Inches(0.8))
        tf = tb.text_frame
        tf.paragraphs[0].text = s
        tf.paragraphs[0].font.size = Pt(14)
        tf.paragraphs[0].font.color.rgb = DARK_GRAY
        y += 1.1

def add_workflow_slide():
    """第7页：完整工作流程"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_GRAY
    bg.line.fill.background()
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.7))
    tf = title_box.text_frame
    tf.paragraphs[0].text = "完整工作流程"
    tf.paragraphs[0].font.size = Pt(36)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BLUE
    
    # 少儿口语工作流
    kids_header = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.2), Inches(5.8), Inches(0.5))
    kids_header.fill.solid()
    kids_header.fill.fore_color.rgb = ORANGE
    kids_header.line.fill.background()
    
    kids_title = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(5.8), Inches(0.5))
    tf = kids_title.text_frame
    tf.paragraphs[0].text = "🎯 少儿口语练习工作流"
    tf.paragraphs[0].font.size = Pt(18)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    kids_steps = [
        "1️⃣ 选择场景：动物园/生日派对/校园",
        "2️⃣ 开始对话：智能体扮演角色互动",
        "3️⃣ 语音输入：孩子自然开口说话",
        "4️⃣ AI识别：分析发音与语义",
        "5️⃣ 智能回应：继续对话并给出鼓励",
        "6️⃣ 勋章奖励：完成对话获得星星"
    ]
    
    y = 2.0
    for step in kids_steps:
        tb = slide.shapes.add_textbox(Inches(0.7), Inches(y), Inches(5.4), Inches(0.45))
        tf = tb.text_frame
        tf.paragraphs[0].text = step
        tf.paragraphs[0].font.size = Pt(13)
        tf.paragraphs[0].font.color.rgb = DARK_GRAY
        y += 0.5
    
    # 三级口译工作流
    int_header = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7), Inches(1.2), Inches(5.8), Inches(0.5))
    int_header.fill.solid()
    int_header.fill.fore_color.rgb = BLUE
    int_header.line.fill.background()
    
    int_title = slide.shapes.add_textbox(Inches(7), Inches(1.2), Inches(5.8), Inches(0.5))
    tf = int_title.text_frame
    tf.paragraphs[0].text = "💼 三级口译训练工作流"
    tf.paragraphs[0].font.size = Pt(18)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    int_steps = [
        "1️⃣ 选择题型：真题/数字/模拟考",
        "2️⃣ 播放音频：设定语速（80-140词/分）",
        "3️⃣ 口译输出：实时录音上传",
        "4️⃣ AI评分：三维度精准评估",
        "5️⃣ 查看报告：错误分析+建议",
        "6️⃣ 继续练习：针对性强化薄弱环节"
    ]
    
    y = 2.0
    for step in int_steps:
        tb = slide.shapes.add_textbox(Inches(7.2), Inches(y), Inches(5.4), Inches(0.45))
        tf = tb.text_frame
        tf.paragraphs[0].text = step
        tf.paragraphs[0].font.size = Pt(13)
        tf.paragraphs[0].font.color.rgb = DARK_GRAY
        y += 0.5
    
    # 协同工作流
    synergy_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(5.2), Inches(12), Inches(1.8))
    synergy_box.fill.solid()
    synergy_box.fill.fore_color.rgb = RGBColor(240, 245, 255)
    synergy_box.line.fill.background()
    
    synergy_title = slide.shapes.add_textbox(Inches(1), Inches(5.3), Inches(11), Inches(0.5))
    tf = synergy_title.text_frame
    tf.paragraphs[0].text = "🔄 全链路学习协同"
    tf.paragraphs[0].font.size = Pt(20)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BLUE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    synergy_flow = slide.shapes.add_textbox(Inches(1), Inches(5.9), Inches(11), Inches(0.9))
    tf = synergy_flow.text_frame
    tf.paragraphs[0].text = "4-12岁：少儿口语启蒙 → 建立自信 → 培养语感 → 18岁+：三级口译进阶 → 专业训练 → 考试通关"
    tf.paragraphs[0].font.size = Pt(16)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = DARK_GRAY
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

def add_comparison_slide():
    """第8页：两者对比与协同"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT_GRAY
    bg.line.fill.background()
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.7))
    tf = title_box.text_frame
    tf.paragraphs[0].text = "同一技术底座，不同能力目标"
    tf.paragraphs[0].font.size = Pt(36)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BLUE
    
    table_x = 0.5
    table_y = 1.2
    table_w = 12
    table_h = 3.5
    
    headers = ["对比项", "少儿口语", "三级口译"]
    header_colors = [BLUE, ORANGE, BLUE]
    for i, header in enumerate(headers):
        w = 4
        header_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(table_x + i * w), Inches(table_y), Inches(w), Inches(0.6))
        header_box.fill.solid()
        header_box.fill.fore_color.rgb = header_colors[i]
        header_box.line.fill.background()
        
        tb = slide.shapes.add_textbox(Inches(table_x + i * w), Inches(table_y), Inches(w), Inches(0.6))
        tf = tb.text_frame
        tf.paragraphs[0].text = header
        tf.paragraphs[0].font.size = Pt(18)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = WHITE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    rows = [
        ["评价标准", "鼓励为主 + 可理解性", "专业评分 + 信息保真"],
        ["反馈方式", "动画 + 语音表情", "文本报告 + 波形对比"],
        ["学习时长", "5–15分钟/次", "20–45分钟/次"],
    ]
    
    row_y = table_y + 0.6
    for row in rows:
        for i, cell in enumerate(row):
            w = 4
            cell_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(table_x + i * w), Inches(row_y), Inches(w), Inches(0.8))
            cell_bg.fill.solid()
            cell_bg.fill.fore_color.rgb = WHITE if (row_y - table_y - 0.6) % 1.6 < 0.8 else RGBColor(240, 244, 248)
            cell_bg.line.fill.background()
            
            tb = slide.shapes.add_textbox(Inches(table_x + i * w + 0.1), Inches(row_y), Inches(w - 0.2), Inches(0.8))
            tf = tb.text_frame
            tf.paragraphs[0].text = cell
            tf.paragraphs[0].font.size = Pt(14)
            tf.paragraphs[0].font.color.rgb = DARK_GRAY
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        row_y += 0.8
    
    synergy_bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(5), Inches(12), Inches(1.5))
    synergy_bg.fill.solid()
    synergy_bg.fill.fore_color.rgb = RGBColor(248, 250, 252)
    synergy_bg.line.fill.background()
    
    synergy_title = slide.shapes.add_textbox(Inches(0.8), Inches(5), Inches(11), Inches(0.4))
    tf = synergy_title.text_frame
    tf.paragraphs[0].text = "✨ 协同可能"
    tf.paragraphs[0].font.size = Pt(20)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = ORANGE
    
    synergy_text = slide.shapes.add_textbox(Inches(0.8), Inches(5.5), Inches(11), Inches(0.8))
    tf = synergy_text.text_frame
    tf.paragraphs[0].text = "少儿智能体积累的发音自信 + 语感 → 未来切入口译训练时焦虑感降低，语音识别自我修正能力更强。"
    tf.paragraphs[0].font.size = Pt(16)
    tf.paragraphs[0].font.color.rgb = DARK_GRAY
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

def add_conclusion_slide():
    """第9页：结语与展望"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    bg.line.fill.background()
    
    circle1 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(2), Inches(1.5), Inches(3), Inches(3))
    circle1.fill.solid()
    circle1.fill.fore_color.rgb = ORANGE
    circle1.line.fill.background()
    
    circle2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.333), Inches(1.5), Inches(3), Inches(3))
    circle2.fill.solid()
    circle2.fill.fore_color.rgb = BLUE
    circle2.line.fill.background()
    
    icon1 = slide.shapes.add_textbox(Inches(2.5), Inches(2), Inches(2), Inches(2))
    tf = icon1.text_frame
    tf.paragraphs[0].text = "🗣️"
    tf.paragraphs[0].font.size = Pt(80)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    icon2 = slide.shapes.add_textbox(Inches(8.833), Inches(2), Inches(2), Inches(2))
    tf = icon2.text_frame
    tf.paragraphs[0].text = "🎧"
    tf.paragraphs[0].font.size = Pt(80)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.8), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    tf.paragraphs[0].text = "让每个阶段都有合适的AI伙伴"
    tf.paragraphs[0].font.size = Pt(36)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BLUE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    summary_box = slide.shapes.add_textbox(Inches(1), Inches(5.6), Inches(11), Inches(0.6))
    tf = summary_box.text_frame
    tf.paragraphs[0].text = "从“说得出”到“译得准”，两个智能体不是替代老师，而是让老师更专注于情感引导与高阶策略教学。"
    tf.paragraphs[0].font.size = Pt(16)
    tf.paragraphs[0].font.color.rgb = DARK_GRAY
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    plans_box = slide.shapes.add_textbox(Inches(1), Inches(6.3), Inches(11), Inches(0.8))
    tf = plans_box.text_frame
    tf.paragraphs[0].text = "Q3：少儿版增加“口译小苗苗”轻量级听力复述模块  |  Q4：三级口译版增加个性化备考路径规划"
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.color.rgb = ORANGE
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    footer_box = slide.shapes.add_textbox(Inches(0.5), Inches(7), Inches(12), Inches(0.5))
    tf = footer_box.text_frame
    tf.paragraphs[0].text = "感谢观看 | 欢迎试用体验  ·  从开口到口译，我们始终陪伴。"
    tf.paragraphs[0].font.size = Pt(18)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BLUE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

# 生成所有页面
add_cover_slide()
add_overview_slide()
add_kids_features_slide()
add_kids_scenarios_slide()
add_interpreter_features_slide()
add_interpreter_scenarios_slide()
add_workflow_slide()
add_comparison_slide()
add_conclusion_slide()

output_path = "AI赋能语言进阶_最终版.pptx"
prs.save(output_path)
print(f"✅ 最终版PPT已生成: {output_path}")
