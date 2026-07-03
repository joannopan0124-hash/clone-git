from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
MEDIUM_GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT_GRAY = RGBColor(0xCC, 0xCC, 0xCC)
OFF_WHITE = RGBColor(0xF5, 0xF5, 0xF5)

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)


def add_text_box(slide, left, top, width, height, text, font_size=18, bold=False, color=BLACK, alignment=PP_ALIGN.LEFT, font_name='Microsoft YaHei'):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_rectangle(slide, left, top, width, height, fill_color=BLACK, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
    else:
        shape.line.fill.background()
    return shape


def add_line(slide, left, top, width, height=Pt(1), color=BLACK):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def create_cover_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    add_rectangle(slide, Inches(0), Inches(0), Inches(6.5), SLIDE_HEIGHT, fill_color=BLACK)
    
    add_rectangle(slide, Inches(7), Inches(2.5), Inches(0.15), Inches(3.5), fill_color=BLACK)
    
    add_text_box(slide, Inches(7.5), Inches(2.3), Inches(5.5), Inches(1.5), 
                 '文档翻译系统', font_size=54, bold=True, color=BLACK)
    
    add_text_box(slide, Inches(7.5), Inches(3.6), Inches(5), Inches(0.8), 
                 'Document Translation System', font_size=20, color=MEDIUM_GRAY)
    
    add_text_box(slide, Inches(7.5), Inches(4.5), Inches(5), Inches(0.6), 
                 '—— 图片翻译 · 实时字幕 · 术语库 · 记忆库', font_size=16, color=MEDIUM_GRAY)
    
    add_text_box(slide, Inches(7.5), Inches(6.5), Inches(5), Inches(0.5), 
                 '2026', font_size=14, color=LIGHT_GRAY)
    
    add_text_box(slide, Inches(0.8), Inches(1), Inches(5), Inches(0.5), 
                 'TRANSLATION', font_size=12, color=WHITE)
    
    add_text_box(slide, Inches(0.8), Inches(6.5), Inches(5), Inches(0.5), 
                 'PAGE 01', font_size=11, color=LIGHT_GRAY)
    
    return slide


def create_toc_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    add_rectangle(slide, Inches(0), Inches(3), SLIDE_WIDTH, Inches(0.02), fill_color=LIGHT_GRAY)
    
    add_text_box(slide, Inches(1), Inches(1), Inches(10), Inches(1), 
                 'CONTENTS', font_size=14, color=MEDIUM_GRAY)
    
    add_text_box(slide, Inches(1), Inches(1.5), Inches(10), Inches(1), 
                 '目  录', font_size=36, bold=True, color=BLACK)
    
    items = [
        ('01', '设计思路', 'Design Concept'),
        ('02', '视觉风格', 'Visual Style'),
        ('03', '图片翻译', 'Image Translation'),
        ('04', '快速翻译', 'Quick Translation'),
        ('05', '实时字幕', 'Real-time Subtitle'),
        ('06', '术语库', 'Glossary'),
        ('07', '记忆库', 'Translation Memory'),
    ]
    
    for i, (num, title, subtitle) in enumerate(items):
        row = i // 4
        col = i % 4
        
        left = Inches(1 + col * 3)
        top = Inches(3.8 + row * 1.6)
        
        add_text_box(slide, left, top, Inches(0.8), Inches(0.5), 
                     num, font_size=28, bold=True, color=BLACK)
        
        add_text_box(slide, left + Inches(0.9), top + Inches(0.05), Inches(2), Inches(0.4), 
                     title, font_size=16, bold=True, color=BLACK)
        
        add_text_box(slide, left + Inches(0.9), top + Inches(0.45), Inches(2), Inches(0.3), 
                     subtitle, font_size=11, color=MEDIUM_GRAY)
    
    add_text_box(slide, Inches(11.5), Inches(7), Inches(1.5), Inches(0.4), 
                 'PAGE 02', font_size=11, color=LIGHT_GRAY, alignment=PP_ALIGN.RIGHT)
    
    return slide


def create_concept_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    add_rectangle(slide, Inches(9.5), Inches(0), Inches(3.833), SLIDE_HEIGHT, fill_color=BLACK)
    
    add_text_box(slide, Inches(10), Inches(1), Inches(3), Inches(0.5), 
                 '01', font_size=14, color=LIGHT_GRAY)
    
    add_text_box(slide, Inches(10), Inches(1.6), Inches(3), Inches(1), 
                 '设计思路', font_size=28, bold=True, color=WHITE)
    
    add_text_box(slide, Inches(10), Inches(2.5), Inches(3), Inches(0.5), 
                 'Design Concept', font_size=14, color=LIGHT_GRAY)
    
    add_text_box(slide, Inches(1), Inches(1), Inches(8), Inches(0.5), 
                 'DESIGN PHILOSOPHY', font_size=12, color=MEDIUM_GRAY)
    
    add_text_box(slide, Inches(1), Inches(1.6), Inches(8), Inches(1.2), 
                 '简洁·高效·专业', font_size=40, bold=True, color=BLACK)
    
    add_line(slide, Inches(1), Inches(3.2), Inches(2), height=Pt(3), color=BLACK)
    
    concepts = [
        ('以用户为中心', '专注于翻译体验的流畅性和直观性，减少操作步骤，提升工作效率'),
        ('专业级工具', '提供术语库、记忆库等专业翻译工具，满足企业级翻译需求'),
        ('多场景适配', '支持图片翻译、文字翻译、实时字幕等多种应用场景'),
    ]
    
    for i, (title, desc) in enumerate(concepts):
        top = Inches(3.8 + i * 1.2)
        
        add_text_box(slide, Inches(1), top, Inches(0.3), Inches(0.5), 
                     '—', font_size=18, bold=True, color=BLACK)
        
        add_text_box(slide, Inches(1.5), top - Inches(0.05), Inches(7), Inches(0.5), 
                     title, font_size=18, bold=True, color=BLACK)
        
        add_text_box(slide, Inches(1.5), top + Inches(0.4), Inches(7), Inches(0.6), 
                     desc, font_size=13, color=MEDIUM_GRAY)
    
    add_text_box(slide, Inches(11.5), Inches(7), Inches(1.5), Inches(0.4), 
                 'PAGE 03', font_size=11, color=LIGHT_GRAY, alignment=PP_ALIGN.RIGHT)
    
    return slide


def create_visual_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    add_rectangle(slide, Inches(0), Inches(0), Inches(5), SLIDE_HEIGHT, fill_color=OFF_WHITE)
    
    add_rectangle(slide, Inches(0.8), Inches(1), Inches(3.5), Inches(5.5), fill_color=WHITE)
    
    add_rectangle(slide, Inches(1.2), Inches(1.4), Inches(2.7), Inches(1.5), fill_color=BLACK)
    add_text_box(slide, Inches(1.4), Inches(1.9), Inches(2.5), Inches(0.5), 
                 'B & W', font_size=24, bold=True, color=WHITE, alignment=PP_ALIGN.CENTER)
    
    add_rectangle(slide, Inches(1.2), Inches(3.1), Inches(1.3), Inches(1.3), fill_color=BLACK)
    add_rectangle(slide, Inches(2.6), Inches(3.1), Inches(1.3), Inches(1.3), fill_color=LIGHT_GRAY)
    add_rectangle(slide, Inches(1.2), Inches(4.5), Inches(1.3), Inches(1.3), fill_color=MEDIUM_GRAY)
    add_rectangle(slide, Inches(2.6), Inches(4.5), Inches(1.3), Inches(1.3), fill_color=DARK_GRAY)
    
    add_text_box(slide, Inches(5.8), Inches(1), Inches(7), Inches(0.5), 
                 '02', font_size=14, color=MEDIUM_GRAY)
    
    add_text_box(slide, Inches(5.8), Inches(1.5), Inches(7), Inches(1), 
                 '视觉风格', font_size=36, bold=True, color=BLACK)
    
    add_text_box(slide, Inches(5.8), Inches(2.5), Inches(7), Inches(0.5), 
                 'Visual Style', font_size=16, color=MEDIUM_GRAY)
    
    add_line(slide, Inches(5.8), Inches(3.2), Inches(1.5), height=Pt(3), color=BLACK)
    
    features = [
        ('黑白商务风', '采用经典黑白灰配色，彰显专业与品质'),
        ('楷体字体', '按钮及标题采用楷体，兼具传统与现代美感'),
        ('扁平化设计', '简洁的界面元素，减少视觉干扰'),
        ('层次分明', '通过字体大小和颜色深浅构建视觉层级'),
    ]
    
    for i, (title, desc) in enumerate(features):
        top = Inches(3.8 + i * 0.9)
        
        add_rectangle(slide, Inches(5.8), top + Inches(0.1), Inches(0.15), Inches(0.5), fill_color=BLACK)
        
        add_text_box(slide, Inches(6.2), top, Inches(6.5), Inches(0.4), 
                     title, font_size=16, bold=True, color=BLACK)
        
        add_text_box(slide, Inches(6.2), top + Inches(0.4), Inches(6.5), Inches(0.4), 
                     desc, font_size=12, color=MEDIUM_GRAY)
    
    add_text_box(slide, Inches(11.5), Inches(7), Inches(1.5), Inches(0.4), 
                 'PAGE 04', font_size=11, color=LIGHT_GRAY, alignment=PP_ALIGN.RIGHT)
    
    return slide


def create_feature_slide(prs, num, title, subtitle, features, left_side=True):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    if left_side:
        add_rectangle(slide, Inches(0), Inches(0), Inches(5), SLIDE_HEIGHT, fill_color=BLACK)
        
        add_text_box(slide, Inches(0.8), Inches(1), Inches(4), Inches(0.5), 
                     num, font_size=14, color=LIGHT_GRAY)
        
        add_text_box(slide, Inches(0.8), Inches(1.6), Inches(4), Inches(1.2), 
                     title, font_size=32, bold=True, color=WHITE)
        
        add_text_box(slide, Inches(0.8), Inches(2.8), Inches(4), Inches(0.5), 
                     subtitle, font_size=14, color=LIGHT_GRAY)
        
        add_line(slide, Inches(0.8), Inches(3.5), Inches(1), height=Pt(2), color=WHITE)
        
        feature_top = 4
        for i, feature in enumerate(features):
            top = Inches(feature_top + i * 0.8)
            add_text_box(slide, Inches(0.8), top, Inches(4), Inches(0.5), 
                         f'— {feature}', font_size=13, color=WHITE)
        
        content_left = Inches(5.8)
    else:
        add_rectangle(slide, Inches(8.333), Inches(0), Inches(5), SLIDE_HEIGHT, fill_color=BLACK)
        
        add_text_box(slide, Inches(9), Inches(1), Inches(4), Inches(0.5), 
                     num, font_size=14, color=LIGHT_GRAY, alignment=PP_ALIGN.RIGHT)
        
        add_text_box(slide, Inches(9), Inches(1.6), Inches(4), Inches(1.2), 
                     title, font_size=32, bold=True, color=WHITE, alignment=PP_ALIGN.RIGHT)
        
        add_text_box(slide, Inches(9), Inches(2.8), Inches(4), Inches(0.5), 
                     subtitle, font_size=14, color=LIGHT_GRAY, alignment=PP_ALIGN.RIGHT)
        
        add_line(slide, Inches(11.5), Inches(3.5), Inches(1), height=Pt(2), color=WHITE)
        
        feature_top = 4
        for i, feature in enumerate(features):
            top = Inches(feature_top + i * 0.8)
            add_text_box(slide, Inches(8.5), top, Inches(4.3), Inches(0.5), 
                         f'{feature} —', font_size=13, color=WHITE, alignment=PP_ALIGN.RIGHT)
        
        content_left = Inches(0.8)
    
    if left_side:
        feature_details = features
        for i, (feat_title, feat_desc) in enumerate([
            (features[0], '上传图片即可自动识别文字并翻译'),
            (features[1] if len(features) > 1 else features[0], '支持多种语言互译'),
        ] if len(features) >= 2 else [(features[0], '详细描述')]):
            pass
    
    if left_side:
        desc_items = features[:3] if len(features) >= 3 else features
        for i, item in enumerate(desc_items):
            top = Inches(2 + i * 1.5)
            
            add_rectangle(slide, content_left, top + Inches(0.1), Inches(0.1), Inches(1), fill_color=BLACK)
            
            add_text_box(slide, content_left + Inches(0.5), top, Inches(6.5), Inches(0.5), 
                         item, font_size=20, bold=True, color=BLACK)
    else:
        desc_items = features[:3] if len(features) >= 3 else features
        for i, item in enumerate(desc_items):
            top = Inches(2 + i * 1.5)
            
            add_rectangle(slide, Inches(8) - Inches(0.1), top + Inches(0.1), Inches(0.1), Inches(1), fill_color=BLACK)
            
            add_text_box(slide, Inches(1.5), top, Inches(6.5), Inches(0.5), 
                         item, font_size=20, bold=True, color=BLACK, alignment=PP_ALIGN.RIGHT)
    
    page_num = int(num)
    add_text_box(slide, Inches(11.5), Inches(7), Inches(1.5), Inches(0.4), 
                 f'PAGE {page_num:02d}', font_size=11, color=LIGHT_GRAY, alignment=PP_ALIGN.RIGHT)
    
    return slide


def create_image_translation_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    add_rectangle(slide, Inches(0), Inches(0), Inches(5), SLIDE_HEIGHT, fill_color=BLACK)
    
    add_text_box(slide, Inches(0.8), Inches(1), Inches(4), Inches(0.5), 
                 '03', font_size=14, color=LIGHT_GRAY)
    
    add_text_box(slide, Inches(0.8), Inches(1.6), Inches(4), Inches(1.2), 
                 '图片翻译', font_size=32, bold=True, color=WHITE)
    
    add_text_box(slide, Inches(0.8), Inches(2.8), Inches(4), Inches(0.5), 
                 'Image Translation', font_size=14, color=LIGHT_GRAY)
    
    add_line(slide, Inches(0.8), Inches(3.5), Inches(1), height=Pt(2), color=WHITE)
    
    left_features = [
        'OCR文字识别',
        '多语言支持',
        '自动翻译',
    ]
    for i, feat in enumerate(left_features):
        top = Inches(4 + i * 0.8)
        add_text_box(slide, Inches(0.8), top, Inches(4), Inches(0.5), 
                     f'— {feat}', font_size=13, color=WHITE)
    
    add_rectangle(slide, Inches(6), Inches(1.2), Inches(6.5), Inches(5.2), fill_color=OFF_WHITE)
    
    add_text_box(slide, Inches(6.5), Inches(1.5), Inches(5), Inches(0.5), 
                 '核心功能', font_size=12, color=MEDIUM_GRAY)
    
    add_text_box(slide, Inches(6.5), Inches(2), Inches(5), Inches(0.8), 
                 'OCR识别 + 智能翻译', font_size=24, bold=True, color=BLACK)
    
    add_line(slide, Inches(6.5), Inches(2.9), Inches(1.5), height=Pt(2), color=BLACK)
    
    items = [
        ('PaddleOCR引擎', '基于百度飞桨的高精度文字识别技术'),
        ('支持图片/PDF', '拖拽上传，一键识别翻译'),
        ('腾讯云翻译', '接入腾讯云机器翻译API，准确高效'),
        ('识别度显示', '实时展示OCR识别置信度'),
        ('结果可编辑', '识别结果支持手动修改调整'),
    ]
    
    for i, (title, desc) in enumerate(items):
        top = Inches(3.3 + i * 0.65)
        
        add_rectangle(slide, Inches(6.5), top + Inches(0.1), Inches(0.1), Inches(0.4), fill_color=BLACK)
        
        add_text_box(slide, Inches(6.9), top, Inches(5.5), Inches(0.35), 
                     title, font_size=14, bold=True, color=BLACK)
        
        add_text_box(slide, Inches(6.9), top + Inches(0.32), Inches(5.5), Inches(0.3), 
                     desc, font_size=11, color=MEDIUM_GRAY)
    
    add_text_box(slide, Inches(11.5), Inches(7), Inches(1.5), Inches(0.4), 
                 'PAGE 05', font_size=11, color=LIGHT_GRAY, alignment=PP_ALIGN.RIGHT)
    
    return slide


def create_quick_translation_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    add_rectangle(slide, Inches(8.333), Inches(0), Inches(5), SLIDE_HEIGHT, fill_color=BLACK)
    
    add_text_box(slide, Inches(9), Inches(1), Inches(4), Inches(0.5), 
                 '04', font_size=14, color=LIGHT_GRAY, alignment=PP_ALIGN.RIGHT)
    
    add_text_box(slide, Inches(9), Inches(1.6), Inches(4), Inches(1.2), 
                 '快速翻译', font_size=32, bold=True, color=WHITE, alignment=PP_ALIGN.RIGHT)
    
    add_text_box(slide, Inches(9), Inches(2.8), Inches(4), Inches(0.5), 
                 'Quick Translation', font_size=14, color=LIGHT_GRAY, alignment=PP_ALIGN.RIGHT)
    
    add_line(slide, Inches(11.5), Inches(3.5), Inches(1), height=Pt(2), color=WHITE)
    
    right_features = [
        '即输即译',
        '多种语言',
        '结果秒出',
    ]
    for i, feat in enumerate(right_features):
        top = Inches(4 + i * 0.8)
        add_text_box(slide, Inches(8.5), top, Inches(4.3), Inches(0.5), 
                     f'{feat} —', font_size=13, color=WHITE, alignment=PP_ALIGN.RIGHT)
    
    add_rectangle(slide, Inches(0.8), Inches(1.2), Inches(6.5), Inches(5.2), fill_color=OFF_WHITE)
    
    add_text_box(slide, Inches(1.3), Inches(1.5), Inches(5), Inches(0.5), 
                 '使用场景', font_size=12, color=MEDIUM_GRAY)
    
    add_text_box(slide, Inches(1.3), Inches(2), Inches(5), Inches(0.8), 
                 '文字快速翻译', font_size=24, bold=True, color=BLACK)
    
    add_line(slide, Inches(1.3), Inches(2.9), Inches(1.5), height=Pt(2), color=BLACK)
    
    items = [
        ('无需上传', '直接输入文字即可翻译，无需上传文件'),
        ('自动检测', '源语言自动检测，简化操作流程'),
        ('中英日韩法德', '支持多种主流语言互译'),
        ('快捷键支持', '回车即可快速翻译'),
        ('实时响应', '翻译结果秒级返回，高效便捷'),
    ]
    
    for i, (title, desc) in enumerate(items):
        top = Inches(3.3 + i * 0.65)
        
        add_rectangle(slide, Inches(1.3), top + Inches(0.1), Inches(0.1), Inches(0.4), fill_color=BLACK)
        
        add_text_box(slide, Inches(1.7), top, Inches(5.5), Inches(0.35), 
                     title, font_size=14, bold=True, color=BLACK)
        
        add_text_box(slide, Inches(1.7), top + Inches(0.32), Inches(5.5), Inches(0.3), 
                     desc, font_size=11, color=MEDIUM_GRAY)
    
    add_text_box(slide, Inches(11.5), Inches(7), Inches(1.5), Inches(0.4), 
                 'PAGE 06', font_size=11, color=LIGHT_GRAY, alignment=PP_ALIGN.RIGHT)
    
    return slide


def create_subtitle_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    add_rectangle(slide, Inches(0), Inches(0), Inches(5), SLIDE_HEIGHT, fill_color=BLACK)
    
    add_text_box(slide, Inches(0.8), Inches(1), Inches(4), Inches(0.5), 
                 '05', font_size=14, color=LIGHT_GRAY)
    
    add_text_box(slide, Inches(0.8), Inches(1.6), Inches(4), Inches(1.2), 
                 '实时字幕', font_size=32, bold=True, color=WHITE)
    
    add_text_box(slide, Inches(0.8), Inches(2.8), Inches(4), Inches(0.5), 
                 'Real-time Subtitle', font_size=14, color=LIGHT_GRAY)
    
    add_line(slide, Inches(0.8), Inches(3.5), Inches(1), height=Pt(2), color=WHITE)
    
    left_features = [
        '语音识别',
        '双语对照',
        '实时更新',
    ]
    for i, feat in enumerate(left_features):
        top = Inches(4 + i * 0.8)
        add_text_box(slide, Inches(0.8), top, Inches(4), Inches(0.5), 
                     f'— {feat}', font_size=13, color=WHITE)
    
    add_rectangle(slide, Inches(6), Inches(1.2), Inches(6.5), Inches(5.2), fill_color=OFF_WHITE)
    
    add_text_box(slide, Inches(6.5), Inches(1.5), Inches(5), Inches(0.5), 
                 '界面设计', font_size=12, color=MEDIUM_GRAY)
    
    add_text_box(slide, Inches(6.5), Inches(2), Inches(5), Inches(0.8), 
                 '双框对照显示', font_size=24, bold=True, color=BLACK)
    
    add_line(slide, Inches(6.5), Inches(2.9), Inches(1.5), height=Pt(2), color=BLACK)
    
    items = [
        ('左右布局', '原文与译文左右并排显示，一目了然'),
        ('独立区域', '原文框灰色背景，译文框蓝色背景区分'),
        ('动画提示', '识别中显示跳动圆点动画效果'),
        ('历史记录', '字幕历史滚动查看，方便回顾'),
        ('一键启停', '开始/停止按钮操作简单便捷'),
    ]
    
    for i, (title, desc) in enumerate(items):
        top = Inches(3.3 + i * 0.65)
        
        add_rectangle(slide, Inches(6.5), top + Inches(0.1), Inches(0.1), Inches(0.4), fill_color=BLACK)
        
        add_text_box(slide, Inches(6.9), top, Inches(5.5), Inches(0.35), 
                     title, font_size=14, bold=True, color=BLACK)
        
        add_text_box(slide, Inches(6.9), top + Inches(0.32), Inches(5.5), Inches(0.3), 
                     desc, font_size=11, color=MEDIUM_GRAY)
    
    add_text_box(slide, Inches(11.5), Inches(7), Inches(1.5), Inches(0.4), 
                 'PAGE 07', font_size=11, color=LIGHT_GRAY, alignment=PP_ALIGN.RIGHT)
    
    return slide


def create_glossary_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    add_rectangle(slide, Inches(8.333), Inches(0), Inches(5), SLIDE_HEIGHT, fill_color=BLACK)
    
    add_text_box(slide, Inches(9), Inches(1), Inches(4), Inches(0.5), 
                 '06', font_size=14, color=LIGHT_GRAY, alignment=PP_ALIGN.RIGHT)
    
    add_text_box(slide, Inches(9), Inches(1.6), Inches(4), Inches(1.2), 
                 '术语库', font_size=32, bold=True, color=WHITE, alignment=PP_ALIGN.RIGHT)
    
    add_text_box(slide, Inches(9), Inches(2.8), Inches(4), Inches(0.5), 
                 'Glossary', font_size=14, color=LIGHT_GRAY, alignment=PP_ALIGN.RIGHT)
    
    add_line(slide, Inches(11.5), Inches(3.5), Inches(1), height=Pt(2), color=WHITE)
    
    right_features = [
        '自定义术语',
        '翻译替换',
        '批量管理',
    ]
    for i, feat in enumerate(right_features):
        top = Inches(4 + i * 0.8)
        add_text_box(slide, Inches(8.5), top, Inches(4.3), Inches(0.5), 
                     f'{feat} —', font_size=13, color=WHITE, alignment=PP_ALIGN.RIGHT)
    
    add_rectangle(slide, Inches(0.8), Inches(1.2), Inches(6.5), Inches(5.2), fill_color=OFF_WHITE)
    
    add_text_box(slide, Inches(1.3), Inches(1.5), Inches(5), Inches(0.5), 
                 '核心价值', font_size=12, color=MEDIUM_GRAY)
    
    add_text_box(slide, Inches(1.3), Inches(2), Inches(5), Inches(0.8), 
                 '确保翻译一致性', font_size=24, bold=True, color=BLACK)
    
    add_line(slide, Inches(1.3), Inches(2.9), Inches(1.5), height=Pt(2), color=BLACK)
    
    items = [
        ('术语管理', '添加、编辑、删除自定义术语对'),
        ('自动替换', '翻译时自动应用术语库替换'),
        ('优先级设置', '支持术语优先级，冲突时优先应用'),
        ('导入导出', '支持JSON格式批量导入导出'),
        ('语言筛选', '按语言方向快速筛选术语'),
    ]
    
    for i, (title, desc) in enumerate(items):
        top = Inches(3.3 + i * 0.65)
        
        add_rectangle(slide, Inches(1.3), top + Inches(0.1), Inches(0.1), Inches(0.4), fill_color=BLACK)
        
        add_text_box(slide, Inches(1.7), top, Inches(5.5), Inches(0.35), 
                     title, font_size=14, bold=True, color=BLACK)
        
        add_text_box(slide, Inches(1.7), top + Inches(0.32), Inches(5.5), Inches(0.3), 
                     desc, font_size=11, color=MEDIUM_GRAY)
    
    add_text_box(slide, Inches(11.5), Inches(7), Inches(1.5), Inches(0.4), 
                 'PAGE 08', font_size=11, color=LIGHT_GRAY, alignment=PP_ALIGN.RIGHT)
    
    return slide


def create_memory_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    add_rectangle(slide, Inches(0), Inches(0), Inches(5), SLIDE_HEIGHT, fill_color=BLACK)
    
    add_text_box(slide, Inches(0.8), Inches(1), Inches(4), Inches(0.5), 
                 '07', font_size=14, color=LIGHT_GRAY)
    
    add_text_box(slide, Inches(0.8), Inches(1.6), Inches(4), Inches(1.2), 
                 '记忆库', font_size=32, bold=True, color=WHITE)
    
    add_text_box(slide, Inches(0.8), Inches(2.8), Inches(4), Inches(0.5), 
                 'Translation Memory', font_size=14, color=LIGHT_GRAY)
    
    add_line(slide, Inches(0.8), Inches(3.5), Inches(1), height=Pt(2), color=WHITE)
    
    left_features = [
        '自动记录',
        '智能匹配',
        '效率提升',
    ]
    for i, feat in enumerate(left_features):
        top = Inches(4 + i * 0.8)
        add_text_box(slide, Inches(0.8), top, Inches(4), Inches(0.5), 
                     f'— {feat}', font_size=13, color=WHITE)
    
    add_rectangle(slide, Inches(6), Inches(1.2), Inches(6.5), Inches(5.2), fill_color=OFF_WHITE)
    
    add_text_box(slide, Inches(6.5), Inches(1.5), Inches(5), Inches(0.5), 
                 '工作原理', font_size=12, color=MEDIUM_GRAY)
    
    add_text_box(slide, Inches(6.5), Inches(2), Inches(5), Inches(0.8), 
                 '越用越聪明', font_size=24, bold=True, color=BLACK)
    
    add_line(slide, Inches(6.5), Inches(2.9), Inches(1.5), height=Pt(2), color=BLACK)
    
    items = [
        ('自动保存', '每次翻译成功后自动存入记忆库'),
        ('相似度检索', '85%相似度即可命中，直接复用结果'),
        ('使用统计', '记录每条翻译的使用次数'),
        ('编辑优化', '支持手动编辑优化翻译结果'),
        ('导入导出', '方便备份和团队共享记忆库'),
    ]
    
    for i, (title, desc) in enumerate(items):
        top = Inches(3.3 + i * 0.65)
        
        add_rectangle(slide, Inches(6.5), top + Inches(0.1), Inches(0.1), Inches(0.4), fill_color=BLACK)
        
        add_text_box(slide, Inches(6.9), top, Inches(5.5), Inches(0.35), 
                     title, font_size=14, bold=True, color=BLACK)
        
        add_text_box(slide, Inches(6.9), top + Inches(0.32), Inches(5.5), Inches(0.3), 
                     desc, font_size=11, color=MEDIUM_GRAY)
    
    add_text_box(slide, Inches(11.5), Inches(7), Inches(1.5), Inches(0.4), 
                 'PAGE 09', font_size=11, color=LIGHT_GRAY, alignment=PP_ALIGN.RIGHT)
    
    return slide


def create_end_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    add_rectangle(slide, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT, fill_color=BLACK)
    
    add_rectangle(slide, Inches(3), Inches(2.8), Inches(0.15), Inches(2), fill_color=WHITE)
    
    add_text_box(slide, Inches(3.5), Inches(2.5), Inches(8), Inches(1.5), 
                 'THANK YOU', font_size=72, bold=True, color=WHITE)
    
    add_text_box(slide, Inches(3.5), Inches(4), Inches(8), Inches(0.8), 
                 '感谢观看', font_size=28, color=LIGHT_GRAY)
    
    add_text_box(slide, Inches(3.5), Inches(5), Inches(8), Inches(0.6), 
                 '文档翻译系统 · 让翻译更高效', font_size=16, color=MEDIUM_GRAY)
    
    add_text_box(slide, Inches(1), Inches(1), Inches(5), Inches(0.5), 
                 'TRANSLATION SYSTEM', font_size=12, color=LIGHT_GRAY)
    
    add_text_box(slide, Inches(11.5), Inches(7), Inches(1.5), Inches(0.4), 
                 'PAGE 10', font_size=11, color=MEDIUM_GRAY, alignment=PP_ALIGN.RIGHT)
    
    return slide


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    
    create_cover_slide(prs)
    create_toc_slide(prs)
    create_concept_slide(prs)
    create_visual_slide(prs)
    create_image_translation_slide(prs)
    create_quick_translation_slide(prs)
    create_subtitle_slide(prs)
    create_glossary_slide(prs)
    create_memory_slide(prs)
    create_end_slide(prs)
    
    output_path = '/workspace/文档翻译系统介绍.pptx'
    prs.save(output_path)
    print(f'PPT已生成: {output_path}')
    print(f'共 {len(prs.slides)} 页')


if __name__ == '__main__':
    main()
